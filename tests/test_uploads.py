import base64
import io
import json
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from web.uploads import UploadError, extract_text, image_thumbnail_data_url, max_upload_bytes


class UploadExtractionTests(unittest.TestCase):
    def test_media_types_have_bounded_size_limits(self) -> None:
        self.assertEqual(max_upload_bytes("notes.py"), 10 * 1024 * 1024)
        self.assertEqual(max_upload_bytes("scan.png"), 20 * 1024 * 1024)
        self.assertEqual(max_upload_bytes("clip.mp4"), 100 * 1024 * 1024)

    def test_extracts_utf8_text_csv_and_json(self) -> None:
        text = extract_text("notes.txt", "한글 메모".encode())
        python = extract_text("script.py", b"print('verified')")
        csv_text = extract_text("data.csv", "name,value\nalpha,7".encode())
        json_text = extract_text("data.json", json.dumps({"name": "alpha"}).encode())

        self.assertEqual(text.text, "한글 메모")
        self.assertFalse(text.truncated)
        self.assertIn("print('verified')", python.text)
        self.assertIn("name | value", csv_text.text)
        self.assertIn('"name": "alpha"', json_text.text)

    def test_extracts_docx_and_xlsx(self) -> None:
        document_stream = io.BytesIO()
        document = Document()
        document.add_paragraph("DOCX evidence")
        document.save(document_stream)
        workbook_stream = io.BytesIO()
        workbook = Workbook()
        workbook.active.append(["metric", 42])
        workbook.save(workbook_stream)

        docx_text = extract_text("report.docx", document_stream.getvalue())
        xlsx_text = extract_text("report.xlsx", workbook_stream.getvalue())

        self.assertEqual(docx_text.text, "DOCX evidence")
        self.assertIn("metric | 42", xlsx_text.text)

    def test_image_is_normalized_for_vision_and_ocr(self) -> None:
        from PIL import Image

        stream = io.BytesIO()
        Image.new("RGB", (100, 60), "white").save(stream, format="PNG")
        with patch("web.uploads.pytesseract.image_to_string", return_value="OCR evidence"):
            result = extract_text("evidence.png", stream.getvalue())

        self.assertIn("OCR evidence", result.text)
        self.assertEqual(len(result.images), 1)
        self.assertEqual(result.images[0][1], "image/jpeg")

    def test_image_thumbnail_is_small_renderable_jpeg(self) -> None:
        from PIL import Image

        source_stream = io.BytesIO()
        Image.new("RGB", (800, 400), "white").save(source_stream, format="JPEG")

        data_url = image_thumbnail_data_url(source_stream.getvalue())
        thumbnail = Image.open(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1])))

        self.assertTrue(data_url.startswith("data:image/jpeg;base64,"))
        self.assertEqual(thumbnail.size, (320, 160))

    def test_video_frames_are_prepared_for_vision_and_ocr(self) -> None:
        from PIL import Image

        frame_stream = io.BytesIO()
        Image.new("RGB", (100, 60), "white").save(frame_stream, format="JPEG")

        def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
            if command[0] == "ffprobe":
                return subprocess.CompletedProcess(command, 0, "12.0\n", "")
            output_pattern = Path(command[-1])
            Path(str(output_pattern).replace("%02d", "01")).write_bytes(frame_stream.getvalue())
            return subprocess.CompletedProcess(command, 0, "", "")

        with patch("web.uploads.subprocess.run", side_effect=fake_run), patch(
            "web.uploads.pytesseract.image_to_string", return_value="frame OCR"
        ):
            result = extract_text("clip.mp4", b"synthetic video")

        self.assertIn("12.0 seconds", result.text)
        self.assertIn("frame OCR", result.text)
        self.assertEqual(len(result.images), 1)

    def test_rejects_empty_and_malformed_documents(self) -> None:
        with self.assertRaisesRegex(UploadError, "empty"):
            extract_text("empty.txt", b"")
        with self.assertRaisesRegex(UploadError, "could not read DOCX"):
            extract_text("broken.docx", b"not a zip file")


if __name__ == "__main__":
    unittest.main()